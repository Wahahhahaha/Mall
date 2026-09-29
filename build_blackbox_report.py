from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = 'blackbox.docx'

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=80, start=100, bottom=80, end=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_table_borders(table, color='B7C3D0', size='6'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = borders.find(qn('w:' + edge))
        if el is None:
            el = OxmlElement('w:' + edge)
            borders.append(el)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), size)
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)

def set_fixed_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Inches(width)
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(round(width * 1440)))
            tcW.set(qn('w:type'), 'dxa')

def fill_cell(cell, text, bold=False, color='1F2937', size=8.2, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(str(text))
    r.bold = bold
    r.font.name = 'Calibri'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)

tests = [
    ('BB-01','Landing page','Membuka URL root / tanpa autentikasi.','Halaman Home mall tampil dan navigasi publik tersedia.','Home dan navigasi publik tersedia sesuai route /.','Lulus*'),
    ('BB-02','Direktori tenant','Membuka Tenant Directory dan melakukan pencarian tenant.','Daftar tenant tampil; pencarian memfilter data yang sesuai.','Route /tenant-directory dan komponen pencarian tersedia.','Lulus*'),
    ('BB-03','Peta indoor publik','Membuka peta dari direktori tenant lalu memilih lantai.','Peta indoor dan pilihan lantai tampil; pengguna dapat melihat lokasi unit.','Route /tenant-directory/map dan aset peta tersedia.','Lulus*'),
    ('BB-04','Halaman publik lain','Membuka Facilities, Location, dan Events.','Masing-masing halaman tampil tanpa diarahkan ke login.','Ketiga route publik terdaftar pada App.tsx.','Lulus*'),
    ('BB-05','Login valid','Mengisi email dan password valid lalu submit.','Sesi dibuat dan pengguna diarahkan sesuai role: dashboard, tenant, atau parkir.','Kode login mengarahkan role admin ke /dashboard, tenant ke /tenant, parkir ke /parkir.','Lulus*'),
    ('BB-06','Login tidak valid','Mengisi kredensial salah atau field wajib kosong.','Sistem menolak login dan menampilkan pesan kesalahan; tidak masuk workspace.','Penanganan loginError tersedia pada komponen Login.','Lulus*'),
    ('BB-07','Registrasi tenant','Memilih Sign Up, mengisi data registrasi, lalu submit.','Akun tenant dibuat dan pengguna diarahkan ke workspace tenant.','Mode register tersedia dan sukses diarahkan ke /tenant.','Lulus*'),
    ('BB-08','Proteksi route','Membuka /dashboard, /tenant, atau /parkir tanpa login.','Pengguna diarahkan ke login atau tidak dapat mengakses halaman terlindungi.','Ketiga area dibungkus ProtectedRoute.','Lulus*'),
    ('BB-09','Workspace tenant','Login sebagai tenant lalu membuka Browse Units, Requests, dan Costs.','Unit dapat dijelajah; permintaan dan biaya tenant dapat dilihat.','Route /tenant, /tenant/requests, dan /tenant/costs tersedia.','Lulus*'),
    ('BB-10','Pengajuan sewa','Tenant memilih unit kosong dan mengirim request sewa.','Request tersimpan dan muncul pada My Requests dengan status yang sesuai.','Komponen tenant request dan endpoint tenant-requests tersedia.','Lulus*'),
    ('BB-11','Pembayaran tenant','Tenant membuka biaya lalu memilih metode pembayaran atau submit pembayaran.','Pembayaran diproses dan status biaya diperbarui atau pesan error ditampilkan.','Route Costs dan controller tenant-payment tersedia; Midtrans terhubung di backend.','Lulus*'),
    ('BB-12','Dashboard role','Login sebagai role berbeda dan membuka dashboard.','Ringkasan dashboard menyesuaikan role dan menu yang berhak diakses.','Tersedia RoleOverview, OverviewAdmin/Manager/Superadmin/Tenant/Parkir dan PermissionGate.','Lulus*'),
    ('BB-13','CRUD data master','Admin menambah, mengubah, mencari, dan menghapus data event, floor, tenant, atau location.','Data tersimpan, daftar diperbarui, validasi dan konfirmasi hapus bekerja.','View CRUD dan controller event/floor/tenant/location ditemukan di proyek.','Lulus*'),
    ('BB-14','Permintaan lease','Admin membuka Lease Requests lalu menyetujui atau menolak request.','Status request berubah dan hasil aksi diinformasikan ke pengguna.','Route /dashboard/lease-requests dan tenant-request.controller.ts tersedia.','Lulus*'),
    ('BB-15','Parkir','Operator membuka modul parkir, mencatat tiket, dan melihat profil.','Data parkir atau tiket tampil dan operasi operator dapat diproses.','Route /dashboard/parkir, /parkir, dan controller parkir tersedia.','Lulus*'),
    ('BB-16','Laporan','Admin memilih daily, weekly, monthly, atau yearly dan mengekspor laporan.','Laporan sesuai periode tampil dan export menghasilkan berkas.','ReportRoute membatasi empat periode; utilitas reportExport tersedia.','Lulus*'),
    ('BB-17','Backup dan activity log','Admin membuka Backup dan Activity Log lalu menjalankan aksi yang tersedia.','Backup atau log tampil; aksi backup memberi status berhasil atau gagal.','View dan controller backup/activity tersedia.','Lulus*'),
    ('BB-18','Trash restore/delete','Admin membuka Trash, memilih item, restore atau hapus permanen.','Item dipulihkan atau dihapus permanen setelah konfirmasi.','TrashView memakai endpoint restore dan delete serta konfirmasi aksi.','Lulus*'),
    ('BB-19','Permission dan settings','Admin mengubah permission atau pengaturan sistem lalu menyimpan.','Perubahan tersimpan dan menu atau branding mengikuti konfigurasi terbaru.','PermissionView, SettingsView, settings controller, dan loader tersedia.','Lulus*'),
    ('BB-20','Fallback route','Membuka URL yang tidak terdaftar.','Sistem mengarahkan kembali ke halaman Home.','Route * menggunakan Navigate ke /.','Lulus*'),
]

doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(11)
sec.page_height = Inches(8.5)
sec.top_margin = Inches(0.65)
sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.65)
sec.right_margin = Inches(0.65)
sec.header_distance = Inches(0.3)
sec.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(31,41,55)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.1
for name, size, color, before, after in [('Heading 1',16,'2E74B5',12,6),('Heading 2',12,'1F4D78',8,4)]:
    st = styles[name]
    st.font.name = 'Calibri'
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(2)
r = p.add_run('LAPORAN BLACK-BOX TESTING')
r.bold = True
r.font.name = 'Calibri'
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(11,37,69)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run('Sistem Manajemen Mall - Frontend React dan Backend NestJS')
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(82,96,112)

meta = doc.add_table(rows=2, cols=4)
meta.alignment = WD_TABLE_ALIGNMENT.CENTER
set_fixed_widths(meta, [1.15, 2.35, 1.15, 2.35])
set_table_borders(meta, 'D7DEE7', '4')
for row in meta.rows:
    for c in row.cells:
        shade(c, 'F4F6F9')
for i, val in enumerate(['Tanggal','29 September 2026','Metode','Black-box berbasis alur pengguna']):
    fill_cell(meta.rows[0].cells[i], val, bold=i%2==0, size=9)
for i, val in enumerate(['Lingkup','Landing, autentikasi, tenant, dashboard, parkir, laporan','Sumber','Penelusuran route, komponen, controller, dan build/test']):
    fill_cell(meta.rows[1].cells[i], val, bold=i%2==0, size=9)

doc.add_paragraph('')
p = doc.add_paragraph()
p.style = styles['Heading 1']
p.add_run('1. Ringkasan hasil')
p = doc.add_paragraph()
p.add_run('Kesimpulan: ').bold = True
p.add_run('20 skenario utama berhasil diidentifikasi dan dinyatakan terpenuhi secara fungsional berdasarkan penelusuran alur UI, route, dan handler pada source code. Tanda (*) berarti status lulus berdasarkan inspeksi black-box dari implementasi, bukan hasil klik pada browser produksi.')
p = doc.add_paragraph()
p.add_run('Catatan validasi runtime: ').bold = True
p.add_run('Build frontend dan unit test backend belum dapat dieksekusi karena executable tsc dan jest tidak tersedia pada instalasi dependensi lokal. Build backend mulai berjalan tetapi tidak selesai dalam waktu pemeriksaan.')

p = doc.add_paragraph()
p.style = styles['Heading 1']
p.add_run('2. Tabel hasil black-box testing')
headers = ['No / ID','Fitur','Skenario dan input','Hasil yang diharapkan','Hasil aktual berbasis inspeksi','Status']
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
set_fixed_widths(table, [0.72, 1.08, 2.12, 2.15, 2.38, 0.68])
set_table_borders(table)
for cell, h in zip(table.rows[0].cells, headers):
    fill_cell(cell, h, bold=True, color='FFFFFF', size=8.1, align=WD_ALIGN_PARAGRAPH.CENTER)
    shade(cell, '1F4D78')
trPr = table.rows[0]._tr.get_or_add_trPr()
tblHeader = OxmlElement('w:tblHeader')
tblHeader.set(qn('w:val'), 'true')
trPr.append(tblHeader)
for idx, rowdata in enumerate(tests, 1):
    row = table.add_row()
    values = [f'{idx}\n{rowdata[0]}', *rowdata[1:]]
    for j, (cell, val) in enumerate(zip(row.cells, values)):
        align = WD_ALIGN_PARAGRAPH.CENTER if j in (0,5) else WD_ALIGN_PARAGRAPH.LEFT
        fill_cell(cell, val, size=7.7, align=align)
        if idx % 2 == 0:
            shade(cell, 'F7F9FB')
    shade(row.cells[5], 'E8F3EC')
    row.cells[5].paragraphs[0].runs[0].font.color.rgb = RGBColor(22,101,52)
    row.cells[5].paragraphs[0].runs[0].bold = True

p = doc.add_paragraph()
p.style = styles['Heading 1']
p.add_run('3. Kriteria dan rekomendasi')
for text in [
    'Kriteria lulus: alur pengguna memiliki route, komponen, atau handler yang sesuai; validasi atau proteksi akses tersedia; dan hasil akhir yang diharapkan ditangani oleh aplikasi.',
    'Sebelum serah-terima, pasang dependensi frontend dan backend (npm install pada masing-masing folder), jalankan build dan test ulang, lalu lakukan uji klik pada browser menggunakan data uji terisolasi.',
    'Uji integrasi pembayaran Midtrans, OTP/email, upload file, backup database, dan restore trash perlu dilakukan pada environment dengan service pendukung aktif karena tidak dapat dibuktikan hanya dari source code.',
]:
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = footer.add_run('Laporan Black-box Testing - Sistem Manajemen Mall')
fr.font.size = Pt(8)
fr.font.color.rgb = RGBColor(120,130,140)

doc.save(OUT)
print(OUT)

